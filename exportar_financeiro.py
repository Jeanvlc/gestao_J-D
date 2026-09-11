# -*- coding: utf-8 -*-
"""
exportar_financeiro.py
Exporta os resultados do painel financeiro para um Excel completo,
com formatacao visual, abas separadas e totais.
"""

import os
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, numbers
from openpyxl.utils import get_column_letter

# Paleta de cores
VERDE = "1B5E20"
VERDE_CLARO = "E8F5E9"
VERDE_MED = "2E7D32"
VERMELHO = "C62828"
VERMELHO_CLARO = "FFEBEE"
LARANJA = "E65100"
LARANJA_CLARO = "FFF3E0"
AMARELO_CLARO = "FFFDE7"
CINZA = "607D8B"
CINZA_CLARO = "ECEFF1"
BRANCO = "FFFFFF"

FMT_MOEDA = 'R$ #.##0,00'
FMT_DATA = 'DD/MM/AAAA'


def _cabecalho(ws, colunas, cor_fundo=VERDE_MED):
    """Escreve uma linha de cabecalho formatada."""
    for col_idx, titulo in enumerate(colunas, start=1):
        c = ws.cell(row=1, column=col_idx, value=titulo)
        c.font = Font(bold=True, color=BRANCO, size=10)
        c.fill = PatternFill("solid", start_color=cor_fundo)
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 20


def _moeda(ws, row, col, valor):
    c = ws.cell(row=row, column=col, value=valor)
    c.number_format = FMT_MOEDA
    return c


def _data_cel(ws, row, col, valor):
    if isinstance(valor, date):
        c = ws.cell(row=row, column=col, value=valor)
        c.number_format = FMT_DATA
    else:
        ws.cell(row=row, column=col, value=str(valor) if valor else "")


def _total_row(ws, row, label, valor, cor_fundo=VERDE_CLARO):
    ws.cell(row=row, column=1, value=label).font = Font(bold=True)
    c = _moeda(ws, row, 2, valor)
    c.font = Font(bold=True)
    for col in range(1, 8):
        ws.cell(row=row, column=col).fill = PatternFill("solid", start_color=cor_fundo)


def exportar_painel(caminho, atrasados, a_vencer, fluxo, data_geracao=None):
    """
    caminho: onde salvar o .xlsx
    atrasados: resultado de analise_financeira.analisar_atrasados()
    a_vencer: resultado de analise_financeira.analisar_a_vencer()
    fluxo: resultado de analise_financeira.calcular_fluxo_caixa()
    """
    os.makedirs(os.path.dirname(caminho) or ".", exist_ok=True)
    data_geracao = data_geracao or date.today()
    wb = Workbook()

    # ------------------------------------------------------------------
    # Aba 1: Resumo
    # ------------------------------------------------------------------
    ws_resumo = wb.active
    ws_resumo.title = "Resumo"
    ws_resumo.column_dimensions["A"].width = 38
    ws_resumo.column_dimensions["B"].width = 20

    titulo = ws_resumo.cell(row=1, column=1,
                             value="PAINEL FINANCEIRO - V. A. RIBAS PREPARO DE SOLO")
    titulo.font = Font(bold=True, size=13, color=VERDE)
    ws_resumo.merge_cells("A1:B1")

    ws_resumo.cell(row=2, column=1, value=f"Gerado em: {data_geracao.strftime('%d/%m/%Y')}")
    ws_resumo.cell(row=2, column=1).font = Font(color=CINZA, italic=True)

    dados_resumo = [
        ("", ""),
        ("SALDO ATUAL DAS CONTAS BANCÁRIAS", fluxo["saldo_inicial"]),
        ("", ""),
        ("BOLETOS ATRASADOS (A PAGAR)", ""),
        ("  Total vencido a pagar", atrasados["total_pagar"]),
        ("  Quantidade de títulos", len(atrasados["pagar"])),
        ("BOLETOS ATRASADOS (A RECEBER)", ""),
        ("  Total vencido a receber", atrasados["total_receber"]),
        ("  Quantidade de títulos", len(atrasados["receber"])),
        ("", ""),
        ("A VENCER NOS PRÓXIMOS 15 DIAS (A PAGAR)", ""),
        ("  Total a pagar", a_vencer["total_pagar"]),
        ("  Quantidade de títulos", len(a_vencer["pagar"])),
        ("A VENCER NOS PRÓXIMOS 15 DIAS (A RECEBER)", ""),
        ("  Total a receber", a_vencer["total_receber"]),
        ("  Quantidade de títulos", len(a_vencer["receber"])),
        ("", ""),
        ("FLUXO DE CAIXA (30 DIAS)", ""),
        ("  Menor saldo projetado",
         fluxo["menor_saldo"]),
        ("  Data do menor saldo",
         fluxo["data_menor_saldo"].strftime("%d/%m/%Y") if fluxo["data_menor_saldo"] else ""),
        ("  Dias com saldo negativo projetado", fluxo["dias_negativos"]),
    ]

    for row_idx, (label, valor) in enumerate(dados_resumo, start=4):
        ws_resumo.cell(row=row_idx, column=1, value=label)
        if isinstance(valor, float):
            c = ws_resumo.cell(row=row_idx, column=2, value=valor)
            c.number_format = FMT_MOEDA
            if valor < 0:
                c.font = Font(color=VERMELHO, bold=True)
            elif label.startswith("  Total vencido a pagar"):
                c.font = Font(color=VERMELHO)
        elif valor != "":
            ws_resumo.cell(row=row_idx, column=2, value=valor)

    # ------------------------------------------------------------------
    # Aba 2: Boletos atrasados
    # ------------------------------------------------------------------
    ws_atr = wb.create_sheet("Atrasados")
    colunas_atr = ["Descrição", "Fornecedor/Cliente", "Vencimento",
                   "Dias Atraso", "Valor", "Categoria", "Centro de Custo", "Tipo"]
    _cabecalho(ws_atr, colunas_atr, cor_fundo=VERMELHO)
    larguras_atr = [36, 30, 14, 12, 16, 22, 20, 10]
    for i, w in enumerate(larguras_atr, 1):
        ws_atr.column_dimensions[get_column_letter(i)].width = w

    row = 2
    for item in atrasados["pagar"]:
        cor = VERMELHO_CLARO if item["gravidade"] == "critico" else (
            LARANJA_CLARO if item["gravidade"] == "alto" else AMARELO_CLARO)
        ws_atr.cell(row=row, column=1, value=item["descricao"])
        ws_atr.cell(row=row, column=2, value=item["nome"])
        _data_cel(ws_atr, row, 3, item["vencimento"])
        c_dias = ws_atr.cell(row=row, column=4, value=item["dias_atraso"])
        c_dias.font = Font(color=VERMELHO, bold=True)
        _moeda(ws_atr, row, 5, item["valor"])
        ws_atr.cell(row=row, column=6, value=item["categoria"])
        ws_atr.cell(row=row, column=7, value=item["centro_custo"])
        ws_atr.cell(row=row, column=8, value="A PAGAR")
        for col in range(1, 9):
            ws_atr.cell(row=row, column=col).fill = PatternFill("solid", start_color=cor)
        row += 1

    for item in atrasados["receber"]:
        ws_atr.cell(row=row, column=1, value=item["descricao"])
        ws_atr.cell(row=row, column=2, value=item["nome"])
        _data_cel(ws_atr, row, 3, item["vencimento"])
        ws_atr.cell(row=row, column=4, value=item["dias_atraso"])
        _moeda(ws_atr, row, 5, item["valor"])
        ws_atr.cell(row=row, column=6, value=item["categoria"])
        ws_atr.cell(row=row, column=7, value=item["centro_custo"])
        ws_atr.cell(row=row, column=8, value="A RECEBER")
        row += 1

    if row > 2:
        _total_row(ws_atr, row + 1, "TOTAL A PAGAR ATRASADO", atrasados["total_pagar"], VERMELHO_CLARO)
        _total_row(ws_atr, row + 2, "TOTAL A RECEBER ATRASADO", atrasados["total_receber"], VERDE_CLARO)

    ws_atr.freeze_panes = "A2"

    # ------------------------------------------------------------------
    # Aba 3: A vencer
    # ------------------------------------------------------------------
    ws_venc = wb.create_sheet("A Vencer (15 dias)")
    colunas_venc = ["Descrição", "Fornecedor/Cliente", "Vencimento",
                    "Dias para Vencer", "Valor", "Categoria", "Centro de Custo", "Tipo"]
    _cabecalho(ws_venc, colunas_venc, cor_fundo=LARANJA)
    for i, w in enumerate([36, 30, 14, 16, 16, 22, 20, 10], 1):
        ws_venc.column_dimensions[get_column_letter(i)].width = w

    row = 2
    for tipo, lista in [("A PAGAR", a_vencer["pagar"]), ("A RECEBER", a_vencer["receber"])]:
        for item in lista:
            cor = AMARELO_CLARO if item["urgencia"] == "normal" else LARANJA_CLARO
            ws_venc.cell(row=row, column=1, value=item["descricao"])
            ws_venc.cell(row=row, column=2, value=item["nome"])
            _data_cel(ws_venc, row, 3, item["vencimento"])
            c_dias = ws_venc.cell(row=row, column=4, value=item["dias_para_vencer"])
            if item["urgencia"] != "normal":
                c_dias.font = Font(color=LARANJA, bold=True)
            _moeda(ws_venc, row, 5, item["valor"])
            ws_venc.cell(row=row, column=6, value=item["categoria"])
            ws_venc.cell(row=row, column=7, value=item["centro_custo"])
            ws_venc.cell(row=row, column=8, value=tipo)
            for col in range(1, 9):
                ws_venc.cell(row=row, column=col).fill = PatternFill("solid", start_color=cor)
            row += 1

    if row > 2:
        _total_row(ws_venc, row + 1, "TOTAL A PAGAR (15 dias)", a_vencer["total_pagar"], LARANJA_CLARO)
        _total_row(ws_venc, row + 2, "TOTAL A RECEBER (15 dias)", a_vencer["total_receber"], VERDE_CLARO)
    ws_venc.freeze_panes = "A2"

    # ------------------------------------------------------------------
    # Aba 4: Fluxo de caixa
    # ------------------------------------------------------------------
    ws_fluxo = wb.create_sheet("Fluxo de Caixa (30 dias)")
    colunas_fluxo = ["Data", "Dia da Semana", "Entradas (R$)",
                     "Saídas (R$)", "Saldo do Dia (R$)", "Situação"]
    _cabecalho(ws_fluxo, colunas_fluxo, cor_fundo=VERDE_MED)
    for i, w in enumerate([14, 16, 18, 18, 20, 14], 1):
        ws_fluxo.column_dimensions[get_column_letter(i)].width = w

    dias_semana = ["Seg", "Ter", "Qua", "Qui", "Sex", "Sáb", "Dom"]
    row = 2

    # Linha de saldo inicial
    ws_fluxo.cell(row=row, column=1, value="SALDO INICIAL")
    c = _moeda(ws_fluxo, row, 5, fluxo["saldo_inicial"])
    c.font = Font(bold=True)
    for col in range(1, 7):
        ws_fluxo.cell(row=row, column=col).fill = PatternFill("solid", start_color=VERDE_CLARO)
    row += 1

    for dia in fluxo["dias"]:
        _data_cel(ws_fluxo, row, 1, dia["data"])
        ws_fluxo.cell(row=row, column=2, value=dias_semana[dia["data"].weekday()])
        _moeda(ws_fluxo, row, 3, dia["entradas"])
        _moeda(ws_fluxo, row, 4, dia["saidas"])
        c_saldo = _moeda(ws_fluxo, row, 5, dia["saldo_dia"])
        situacao = "⚠ NEGATIVO" if dia["saldo_negativo"] else "OK"
        ws_fluxo.cell(row=row, column=6, value=situacao)

        if dia["saldo_negativo"]:
            c_saldo.font = Font(color=VERMELHO, bold=True)
            cor = VERMELHO_CLARO
        elif dia["entradas"] > 0 or dia["saidas"] > 0:
            cor = VERDE_CLARO if dia["entradas"] >= dia["saidas"] else AMARELO_CLARO
        else:
            cor = BRANCO

        for col in range(1, 7):
            ws_fluxo.cell(row=row, column=col).fill = PatternFill("solid", start_color=cor)
        row += 1

    ws_fluxo.freeze_panes = "A3"

    wb.save(caminho)
    return caminho
